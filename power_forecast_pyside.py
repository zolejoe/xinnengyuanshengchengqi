import os
import re
import shutil
import sys
from PySide6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
                               QHBoxLayout, QLabel, QLineEdit, QPushButton,
                               QProgressBar, QTextEdit, QFileDialog, QMessageBox)
from PySide6.QtCore import QThread, Signal, Qt
import pandas as pd
from openpyxl import load_workbook

# -------------------- 后台处理线程 --------------------
class WorkerThread(QThread):
    log_signal = Signal(str)
    finished_signal = Signal(bool, str, int)  # success, message, file_count

    def __init__(self, source_file, template_file, output_dir):
        super().__init__()
        self.source_file = source_file
        self.template_file = template_file
        self.output_dir = output_dir

    def log(self, msg):
        self.log_signal.emit(msg)

    def run(self):
        try:
            # 读取源数据
            df = pd.read_excel(self.source_file, sheet_name="总负荷 (营销调整)", header=0)
            time_col = df.columns[0]
            time_data = df[time_col].values
            station_cols = [col for col in df.columns if col != time_col][3:-1]
            if not station_cols:
                raise ValueError("未识别到场站列，请检查源文件格式")

            self.log(f"共找到 {len(station_cols)} 个场站，开始处理...")

            for idx, col in enumerate(station_cols, 1):
                match = re.match(r"(.+?)\s*\(.*\)", col)
                station_name = match.group(1) if match else col
                self.log(f"[{idx}/{len(station_cols)}] 处理场站: {station_name}")

                raw_name = f"发电侧日前交易模版-新能源曲线{station_name}.xlsx"
                clean_name = re.sub(r'[\\/*?:"<>|]', '_', raw_name)
                output_path = os.path.join(self.output_dir, clean_name)

                shutil.copy2(self.template_file, output_path)
                wb = load_workbook(output_path)
                ws = wb.active
                power_data = df[col].fillna(0).values
                start_row = 3
                for i, (t, p) in enumerate(zip(time_data, power_data)):
                    ws.cell(row=start_row + i, column=1, value=t)
                    ws.cell(row=start_row + i, column=2, value=p)
                wb.save(output_path)

            self.log("所有场站处理完成！")
            self.finished_signal.emit(True, f"处理成功！共生成 {len(station_cols)} 个文件。", len(station_cols))
        except Exception as e:
            self.log(f"处理出错: {e}")
            self.finished_signal.emit(False, str(e), 0)


# -------------------- 主窗口 --------------------
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("新能源曲线生成器")
        self.setMinimumSize(600, 400)

        # 变量
        self.source_file = ""
        self.template_file = ""
        self.output_dir = ""

        # 创建控件
        self.source_line = QLineEdit()
        self.template_line = QLineEdit()
        self.output_line = QLineEdit()
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 0)  # 设置为不确定进度条
        self.progress_bar.hide()

        self.run_btn = QPushButton("开始处理")
        self.run_btn.clicked.connect(self.start_processing)

        # 布局
        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout(central)

        # 源文件行
        src_layout = QHBoxLayout()
        src_layout.addWidget(QLabel("功率预测文件："))
        src_layout.addWidget(self.source_line)
        src_btn = QPushButton("浏览")
        src_btn.clicked.connect(lambda: self.browse_file(self.source_line, "选择功率预测文件", "Excel files (*.xlsx *.xls)"))
        src_layout.addWidget(src_btn)
        main_layout.addLayout(src_layout)

        # 模板文件行
        tpl_layout = QHBoxLayout()
        tpl_layout.addWidget(QLabel("模板文件："))
        tpl_layout.addWidget(self.template_line)
        tpl_btn = QPushButton("浏览")
        tpl_btn.clicked.connect(lambda: self.browse_file(self.template_line, "选择模板文件", "Excel files (*.xlsx *.xls)"))
        tpl_layout.addWidget(tpl_btn)
        main_layout.addLayout(tpl_layout)

        # 输出目录行
        out_layout = QHBoxLayout()
        out_layout.addWidget(QLabel("输出目录："))
        out_layout.addWidget(self.output_line)
        out_btn = QPushButton("浏览")
        out_btn.clicked.connect(lambda: self.browse_dir(self.output_line, "选择输出目录"))
        out_layout.addWidget(out_btn)
        main_layout.addLayout(out_layout)

        # 进度条
        main_layout.addWidget(self.progress_bar)

        # 日志
        main_layout.addWidget(QLabel("处理日志："))
        main_layout.addWidget(self.log_text)

        # 按钮
        main_layout.addWidget(self.run_btn, alignment=Qt.AlignCenter)

    def browse_file(self, line_edit, title, file_filter):
        path, _ = QFileDialog.getOpenFileName(self, title, "", file_filter)
        if path:
            line_edit.setText(path)
            self.log(f"已选择: {path}")

    def browse_dir(self, line_edit, title):
        path = QFileDialog.getExistingDirectory(self, title)
        if path:
            line_edit.setText(path)
            self.log(f"已选择输出目录: {path}")

    def log(self, msg):
        self.log_text.append(msg)

    def start_processing(self):
        self.source_file = self.source_line.text().strip()
        self.template_file = self.template_line.text().strip()
        self.output_dir = self.output_line.text().strip()

        if not self.source_file:
            QMessageBox.warning(self, "提示", "请选择功率预测文件")
            return
        if not self.template_file:
            QMessageBox.warning(self, "提示", "请选择模板文件")
            return
        if not self.output_dir:
            QMessageBox.warning(self, "提示", "请选择输出目录")
            return
        if not os.path.exists(self.source_file):
            QMessageBox.critical(self, "错误", "功率预测文件不存在")
            return
        if not os.path.exists(self.template_file):
            QMessageBox.critical(self, "错误", "模板文件不存在")
            return

        os.makedirs(self.output_dir, exist_ok=True)

        self.run_btn.setEnabled(False)
        self.progress_bar.show()
        self.log_text.clear()

        self.worker = WorkerThread(self.source_file, self.template_file, self.output_dir)
        self.worker.log_signal.connect(self.log)
        self.worker.finished_signal.connect(self.on_finished)
        self.worker.start()

    def on_finished(self, success, message, file_count):
        self.progress_bar.hide()
        self.run_btn.setEnabled(True)
        if success:
            QMessageBox.information(self, "完成", message)
        else:
            QMessageBox.critical(self, "错误", message)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())